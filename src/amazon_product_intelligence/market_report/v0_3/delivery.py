"""Artifact Tool delivery of the same localized business-report projection."""

from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path
import shutil
import subprocess
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from amazon_product_intelligence.contracts import canonical_json
from .presentation import markdown, workbook_spec


def workbook_semantics(path):
    """Read-only round-trip check, independent of the XLSX authoring engine."""
    wb = load_workbook(path, data_only=False)
    try:
        return [{"name": s.title, "rows": [[c.value for c in r] for r in s.iter_rows()], "formats": [[c.number_format for c in r] for r in s.iter_rows()], "tables": sorted(t.ref for t in s.tables.values()), "validations": sorted((str(v.sqref), v.type, v.formula1) for v in s.data_validations.dataValidation), "freeze": s.freeze_panes} for s in wb]
    finally:
        wb.close()


def validate_workbook(path, spec):
    wb = load_workbook(path, data_only=False)
    try:
        if wb.sheetnames != [s["name"] for s in spec["sheets"]]:
            raise ValueError("workbook sheet order differs from presentation contract")
        human = wb[spec["sheets"][13]["name"]]
        expected = spec["sheets"][13]["tables"][0]["rows"]
        found = [row for row in human.iter_rows(values_only=True) if row[0] in {r[0] for r in expected}]
        if len(found) != 11 or any(r[1] is not None or r[2] is not None for r in found):
            raise ValueError("eleven human verdict/notes pairs must be blank")
        errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
        if any(c.data_type == "f" or c.data_type == "e" or (isinstance(c.value, str) and c.value in errors) for s in wb for row in s for c in row):
            raise ValueError("snapshot workbook contains unexpected formulas/errors")
        for source, target in zip(spec["sheets"], wb):
            # Compare every exported table cell, including typed numbers, against
            # the report projection. Human inputs alone are deliberately blank.
            native = list(target.tables.values())
            if len(native) != len(source["tables"]):
                raise ValueError("table count differs from presentation contract")
            for table, exported in zip(source["tables"], native):
                actual = [[c.value for c in row] for row in target[exported.ref]]
                wanted = [table["headers"], *table["rows"]]
                if len(actual) != len(wanted):
                    raise ValueError("exported row count differs")
                for left, right in zip(actual, wanted):
                    for a, b in zip(left, right):
                        if b == "":
                            b = None
                        if isinstance(a, str) and a.startswith("'") and isinstance(b, str) and b[:1] in "=+@-":
                            a = a[1:]
                        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                            if abs(a - b) > 1e-10 * max(1, abs(b)):
                                raise ValueError("numeric export mismatch")
                        elif a != b:
                            raise ValueError("exported cell differs from localized projection")
    finally:
        wb.close()
    return {"sheets": 15, "blank_human_cells": 22, "all_table_cells_match": True, "formula_errors": 0}


def deliver(report, output_directory, *, node_executable, renderer_runtime, locale="zh-CN", qa_directory=None):
    """The runtime directory must have access to @oai/artifact-tool.

    Caller supplies runtime paths; machine-local dependencies are not embedded
    in source. The renderer is copied unchanged to a temporary runtime child.
    """
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    spec = workbook_spec(report, locale)
    suffix = locale.lower().replace("-", "_")
    renderer = Path(__file__).with_name("renderer.mjs")
    with TemporaryDirectory(prefix="operator-v03-", dir=renderer_runtime) as directory:
        stage = Path(directory)
        shutil.copyfile(renderer, stage / "renderer.mjs")
        (stage / "spec.json").write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
        xlsx = stage / "report.xlsx"
        command = [str(node_executable), str(stage / "renderer.mjs"), str(stage / "spec.json"), str(xlsx)]
        if qa_directory:
            command.append(str(qa_directory))
        subprocess.run(command, check=True, timeout=600, capture_output=True, text=True, encoding="utf-8")
        checks = validate_workbook(xlsx, spec)
        final_xlsx = output / f"operator_market_report_{suffix}.xlsx"
        shutil.copyfile(xlsx, final_xlsx)
    artifacts = {"market_report_v0_3.json": report.to_json(), f"operator_market_report_{suffix}.md": markdown(report, locale)}
    for name, body in artifacts.items():
        (output / name).write_text(body, encoding="utf-8", newline="\n")
    semantics = workbook_semantics(final_xlsx)
    return {"checks": checks, "json_sha256": sha256((output / "market_report_v0_3.json").read_bytes()).hexdigest(), "markdown_sha256": sha256((output / f"operator_market_report_{suffix}.md").read_bytes()).hexdigest(), "workbook_content_sha256": sha256(canonical_json(semantics).encode()).hexdigest(), "workbook_raw_sha256": sha256(final_xlsx.read_bytes()).hexdigest(), "spec_sha256": sha256(canonical_json(spec).encode()).hexdigest()}

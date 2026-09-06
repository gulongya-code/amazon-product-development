"""Deterministic JSON, Markdown, and XLSX delivery for integrated reports."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from amazon_product_intelligence.contracts import canonical_json
from amazon_product_intelligence.xlsx_delivery.package_fingerprint import (
    ooxml_package_content_sha256,
)

from .models import IntegratedMarketReport


INTEGRATED_MARKET_REPORT_DELIVERY_VERSION = (
    "integrated-route-market-report-delivery-v1.0"
)
MARKET_REPORT_JSON_FILENAME = "market_report.json"
OPERATOR_MARKDOWN_FILENAME = "operator_market_report.md"
OPERATOR_XLSX_FILENAME = "operator_market_report.xlsx"
BUSINESS_VALUE_REVIEW_FILENAME = "business_value_review.md"

_NAVY = "17365D"
_BLUE = "2F75B5"
_PALE_BLUE = "D9EAF7"
_PALE_GREEN = "E2F0D9"
_PALE_AMBER = "FFF2CC"
_PALE_RED = "FCE4D6"
_WHITE = "FFFFFF"
_GREY = "666666"
_THIN = Side(style="thin", color="D9E1F2")
_BORDER = Border(bottom=_THIN)


@dataclass(frozen=True, slots=True, kw_only=True)
class IntegratedMarketReportDeliveryResult:
    source_report_id: str
    delivery_version: str
    json_path: Path
    markdown_path: Path
    xlsx_path: Path
    business_value_review_path: Path
    json_sha256: str
    markdown_sha256: str
    xlsx_sha256: str
    xlsx_package_content_sha256: str
    xlsx_semantic_sha256: str
    business_value_review_sha256: str


def _text(value: Any) -> str:
    if value is None:
        return "UNAVAILABLE"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return str(value).replace("|", "\\|").replace("\n", " ")


def _metric(metric: Mapping[str, Any], *, nested: str | None = None) -> Any:
    if metric.get("availability") == "UNAVAILABLE":
        return None
    value = metric.get("value")
    if nested is not None and isinstance(value, Mapping):
        value = value.get(nested)
    return value


def _fmt(value: Any, *, percentage: bool = False, currency: bool = False) -> str:
    if value is None:
        return "UNAVAILABLE"
    if type(value) not in {int, float}:
        return _text(value)
    if percentage:
        return f"{value:.2%}"
    if currency:
        return f"${value:,.2f}"
    return f"{value:,.2f}" if isinstance(value, float) else f"{value:,}"


def _route_table(report: IntegratedMarketReport) -> list[list[str]]:
    rows = [[
        "Route", "Products", "Share of accepted", "Monthly sales", "Sales share",
        "Median price", "Median reviews", "Top brand share", "Demand efficiency",
        "Opportunity signal",
    ]]
    recommendation = {
        item["route_id"]: item for item in report.opportunity["recommendations"]
    }
    for item in report.route_landscape:
        signal = recommendation[item["route_id"]]
        rows.append([
            item["route_label"],
            _fmt(item["products"]),
            _fmt(item["share_of_accepted"], percentage=True),
            _fmt(_metric(item["monthly_sales"])),
            _fmt(_metric(item["sales_share"]), percentage=True),
            _fmt(_metric(item["price_distribution"], nested="median"), currency=True),
            _fmt(_metric(item["review_count_distribution"], nested="median")),
            _fmt(_metric(item["brand_listing_concentration"], nested="top_1_share"), percentage=True),
            _fmt(_metric(item["demand_efficiency"])),
            signal["classification"],
        ])
    return rows


def render_operator_markdown(report: IntegratedMarketReport) -> str:
    overview = report.market_overview
    counts = overview["coverage_counts"]
    rates = overview["coverage_rates"]
    lines = [
        f"# {report.category} Market Intelligence Report",
        "",
        f"Marketplace: **{report.marketplace}**  ",
        f"Evidence snapshot: **{report.generated_at}**  ",
        f"Report ID: `{report.report_id}`",
        "",
        "> Decision support only. Route recommendations prioritize further research; they do not assert launch viability or profitability.",
        "",
        "## 1. Market overview",
        "",
        "| Measure | Value |",
        "|---|---:|",
        f"| Accepted listings | {_fmt(counts['accepted'])} |",
        f"| Primary cohort | {_fmt(counts['primary_cohort'])} |",
        f"| Assigned | {_fmt(counts['assigned'])} ({_fmt(rates['assigned'], percentage=True)}) |",
        f"| Unclassified | {_fmt(counts['unclassified'])} ({_fmt(rates['unclassified'], percentage=True)}) |",
        f"| Review required | {_fmt(counts['review_required'])} ({_fmt(rates['review_required'], percentage=True)}) |",
        f"| Routes / candidates | {_fmt(overview['route_count'])} / {_fmt(overview['candidate_count'])} |",
        f"| Monthly sales estimate | {_fmt(_metric(overview['monthly_sales']))} |",
        f"| Monthly revenue estimate | {_fmt(_metric(overview['monthly_revenue']), currency=True)} |",
        f"| Median price | {_fmt(_metric(overview['price_distribution'], nested='median'), currency=True)} |",
        f"| Median reviews | {_fmt(_metric(overview['review_count_distribution'], nested='median'))} |",
        "",
        "Important limitations: " + "; ".join(overview["data_limitations"]) + ".",
        "",
        "## 2. Route landscape",
        "",
    ]
    route_table = _route_table(report)
    lines.extend([
        "| " + " | ".join(route_table[0]) + " |",
        "|" + "|".join("---" if index == 0 else "---:" for index in range(len(route_table[0]))) + "|",
        *("| " + " | ".join(row) + " |" for row in route_table[1:]),
        "",
        "## 3. Route comparison",
        "",
        "| Material difference | Leading Route | Observed value | Evidence basis |",
        "|---|---|---:|---|",
    ])
    for item in report.route_comparison:
        lines.append(
            f"| {_text(item['comparison'])} | {_text(item['route_label'])} | {_fmt(item['value'])} | {_text(item['basis'])} |"
        )
    lines.extend([
        "",
        "## 4. Buyer needs",
        "",
        f"Status: **{report.buyer_needs['availability']}**.",
        "",
        *[f"- {item}" for item in report.buyer_needs["limitations"]],
        "",
        "## 5. Competition",
        "",
        "Competition is interpreted at Route level using review barriers and governed brand/seller concentration. "
        "Only observed fields are populated; unavailable values remain unavailable.",
        "",
        "| Route | Median reviews | Top brand listing share | Top brand sales share | Top seller sales share |",
        "|---|---:|---:|---:|---:|",
    ])
    for item in report.route_landscape:
        lines.append(
            "| " + " | ".join([
                _text(item["route_label"]),
                _fmt(_metric(item["review_count_distribution"], nested="median")),
                _fmt(_metric(item["brand_listing_concentration"], nested="top_1_share"), percentage=True),
                _fmt(_metric(item["brand_sales_concentration"], nested="top_1_share"), percentage=True),
                _fmt(_metric(item["seller_sales_concentration"], nested="top_1_share"), percentage=True),
            ]) + " |"
        )
    lines.extend([
        "",
        "## 6. Opportunity",
        "",
        "| Priority | Classification | Route | Quantitative / governed basis |",
        "|---:|---|---|---|",
    ])
    for item in report.opportunity["recommendations"]:
        lines.append(
            f"| {_text(item['priority'])} | {_text(item['classification'])} | {_text(item['route_label'])} | {_text(item['basis'])} |"
        )
    lines.extend([
        "",
        "## 7. Risks / unknowns",
        "",
        "| Risk | Value | Basis |",
        "|---|---:|---|",
        *(
            f"| {_text(item['risk'])} | {_fmt(item['value'], percentage=item['risk'].endswith('_EXPOSURE'))} | {_text(item['basis'])} |"
            for item in report.risks_unknowns
        ),
        "",
        "## Business Value Review (human only)",
        "",
        "| Question | Human verdict | Notes |",
        "|---|---|---|",
        *(f"| {_text(item['question'])} |  |  |" for item in report.business_value_review),
        "",
        "## Audit / provenance",
        "",
        "| Field | Value |",
        "|---|---|",
        *(f"| {_text(key)} | {_text(value)} |" for key, value in sorted(report.source_lineage.items())),
        "",
    ])
    return "\n".join(lines)


def render_business_value_review(report: IntegratedMarketReport) -> str:
    lines = [
        f"# {report.category} Business Value Review",
        "",
        f"Report ID: `{report.report_id}`",
        "",
        "Human verdicts are intentionally blank. Do not infer them from technical acceptance.",
        "",
        "| Question | Human verdict | Notes |",
        "|---|---|---|",
    ]
    lines.extend(
        f"| {_text(item['question'])} |  |  |"
        for item in report.business_value_review
    )
    lines.append("")
    return "\n".join(lines)


def _safe(value: Any) -> Any:
    if value is None:
        return "UNAVAILABLE"
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return value


def _title(sheet: Any, title: str, subtitle: str, columns: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    sheet.cell(1, 1, title)
    sheet.cell(1, 1).font = Font(size=18, bold=True, color=_WHITE)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=_NAVY)
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet.cell(2, 1, subtitle)
    sheet.cell(2, 1).font = Font(italic=True, color=_GREY)
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[2].height = 30


def _table_sheet(
    workbook: Workbook,
    *,
    name: str,
    title: str,
    subtitle: str,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
) -> Any:
    sheet = workbook.create_sheet(name)
    _title(sheet, title, subtitle, len(headers))
    sheet.append([])
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe(value) for value in row])
    header_row = 4
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[header_row].height = 30
    if sheet.max_row > header_row:
        table = Table(displayName=f"T_{name.replace(' ', '_')}", ref=f"A{header_row}:{sheet.cell(header_row, len(headers)).column_letter}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(header_row, len(headers)).column_letter}{sheet.max_row}"
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = _BORDER
    for index, header in enumerate(headers, start=1):
        width = 18
        if any(word in header.lower() for word in ("route", "basis", "limitation", "question", "definition", "value")):
            width = 34
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def render_operator_xlsx(report: IntegratedMarketReport, destination: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "amazon-product-intelligence"
    workbook.properties.lastModifiedBy = "amazon-product-intelligence"
    workbook.properties.created = datetime(2000, 1, 1)
    workbook.properties.modified = datetime(2000, 1, 1)
    overview = report.market_overview
    counts = overview["coverage_counts"]
    rates = overview["coverage_rates"]

    executive = _table_sheet(
        workbook,
        name="Executive Summary",
        title=f"{report.category} Market Intelligence",
        subtitle="Decision support only. Research priorities are evidence-backed, not launch recommendations.",
        headers=["Decision area", "Observed result", "Operator interpretation"],
        rows=[
            ["Market coverage", f"{counts['assigned']}/{counts['accepted']} assigned ({rates['assigned']:.2%})", "Route metrics exclude unresolved and unclassified listings."],
            ["Route structure", f"{overview['route_count']} Routes; {overview['candidate_count']} candidates", "Compare meaningful purchase/product alternatives before deeper validation."],
            ["Priority research", ", ".join(item["route_label"] for item in report.opportunity["recommendations"] if item["classification"] == "PRIORITY_RESEARCH") or "UNAVAILABLE", "Candidate policy identifies investigation order, not business viability."],
            ["Buyer needs", report.buyer_needs["availability"], "; ".join(report.buyer_needs["limitations"])],
            ["Business value", "PENDING HUMAN REVIEW", "Complete the Business Value Review sheet; technical PASS does not fill this gate."],
        ],
    )
    executive.sheet_view.showGridLines = False

    overview_sheet = _table_sheet(
        workbook,
        name="Market Overview",
        title="Market Overview",
        subtitle="Missing observations remain unavailable and are never converted to zero.",
        headers=["Measure", "Value", "Availability / basis"],
        rows=[
            ["Accepted listings", counts["accepted"], "governed dataset"],
            ["Primary cohort", counts["primary_cohort"], "released Semantic/Route result"],
            ["Assigned", counts["assigned"], f"{rates['assigned']:.2%}"],
            ["Unclassified", counts["unclassified"], f"{rates['unclassified']:.2%}"],
            ["Review required", counts["review_required"], f"{rates['review_required']:.2%}"],
            ["Monthly sales", _metric(overview["monthly_sales"]), overview["monthly_sales"]["availability"]],
            ["Monthly revenue USD", _metric(overview["monthly_revenue"]), overview["monthly_revenue"]["availability"]],
            ["Median price USD", _metric(overview["price_distribution"], nested="median"), overview["price_distribution"]["availability"]],
            ["Median reviews", _metric(overview["review_count_distribution"], nested="median"), overview["review_count_distribution"]["availability"]],
            ["Top brand listing share", overview["brand_concentration"].get("top_1_share"), overview["brand_concentration"]["availability"]],
            ["Top seller listing share", overview["seller_concentration"].get("top_1_share"), overview["seller_concentration"]["availability"]],
        ],
    )
    for row in (5, 6, 7, 8, 9, 10, 13):
        overview_sheet.cell(row, 2).number_format = "#,##0"
    overview_sheet.cell(11, 2).number_format = '"$"#,##0'
    overview_sheet.cell(12, 2).number_format = '"$"#,##0.00'
    overview_sheet.cell(13, 2).number_format = "#,##0"
    for row in (14, 15):
        overview_sheet.cell(row, 2).number_format = "0.0%"

    route_headers = [
        "Route", "Definition", "Products", "Share accepted", "Monthly sales",
        "Sales share", "Monthly revenue USD", "Median price USD", "Median reviews",
        "Median rating", "Top brand listing share", "Top seller sales share",
        "Demand efficiency", "New-product efficiency", "MoM growth", "YoY growth",
        "Opportunity signal", "Evidence / limitations",
    ]
    recommendations = {item["route_id"]: item for item in report.opportunity["recommendations"]}
    route_rows = []
    for item in report.route_landscape:
        route_rows.append([
            item["route_label"], item["route_definition"], item["products"], item["share_of_accepted"],
            _metric(item["monthly_sales"]), _metric(item["sales_share"]), _metric(item["monthly_revenue"]),
            _metric(item["price_distribution"], nested="median"),
            _metric(item["review_count_distribution"], nested="median"),
            _metric(item["rating_distribution"], nested="median"),
            _metric(item["brand_listing_concentration"], nested="top_1_share"),
            _metric(item["seller_sales_concentration"], nested="top_1_share"),
            _metric(item["demand_efficiency"]), _metric(item["new_product_demand_efficiency"]),
            _metric(item["mom_growth"], nested="aggregate_growth"),
            _metric(item["yoy_growth"], nested="aggregate_growth"),
            recommendations[item["route_id"]]["classification"],
            (
                f"title {item['evidence_coverage']['title']:.0%}; "
                f"structured {item['evidence_coverage']['structured_parameters']:.0%}; "
                f"{len(item['assignment_limitations'])} assignment flags"
            ),
        ])
    route_sheet = _table_sheet(
        workbook, name="Route Landscape", title="Route Landscape",
        subtitle="First-class Route comparison. Monetary and demand values use only known observations.",
        headers=route_headers, rows=route_rows,
    )
    for column in (4, 6, 11, 12, 15, 16):
        for cell in route_sheet.iter_cols(min_col=column, max_col=column, min_row=5):
            for item in cell:
                item.number_format = "0.00%"
    for column in (7, 8):
        for cell in route_sheet.iter_cols(min_col=column, max_col=column, min_row=5):
            for item in cell:
                item.number_format = '$#,##0.00'
    for column in (3, 5, 9):
        for cell in route_sheet.iter_cols(min_col=column, max_col=column, min_row=5):
            for item in cell:
                item.number_format = "#,##0"
    for column in (10, 13, 14):
        for cell in route_sheet.iter_cols(min_col=column, max_col=column, min_row=5):
            for item in cell:
                item.number_format = "0.00"

    _table_sheet(
        workbook, name="Route Comparison", title="Material Route Differences",
        subtitle="Each comparison identifies the source metric used; unavailable comparisons remain explicit.",
        headers=["Comparison", "Availability", "Leading Route", "Observed value", "Basis"],
        rows=[[
            item["comparison"].replace("_", " ").title(), item["availability"],
            item["route_label"], item["value"], item["basis"],
        ] for item in report.route_comparison],
    )
    _table_sheet(
        workbook, name="Buyer Needs", title="Buyer Needs",
        subtitle="No Buyer Need is inferred from listing marketing text.",
        headers=["Availability", "Theme / link", "Evidence basis or limitation"],
        rows=[[report.buyer_needs["availability"], "UNAVAILABLE", limitation] for limitation in report.buyer_needs["limitations"]],
    )
    competition_sheet = _table_sheet(
        workbook, name="Competition", title="Route Competition",
        subtitle="Review barriers and brand/seller concentration are descriptive competition indicators.",
        headers=["Route", "Median reviews", "Top brand listing share", "Top brand sales share", "Top seller listing share", "Top seller sales share"],
        rows=[[
            item["route_label"], _metric(item["review_count_distribution"], nested="median"),
            _metric(item["brand_listing_concentration"], nested="top_1_share"),
            _metric(item["brand_sales_concentration"], nested="top_1_share"),
            _metric(item["seller_listing_concentration"], nested="top_1_share"),
            _metric(item["seller_sales_concentration"], nested="top_1_share"),
        ] for item in report.route_landscape],
    )
    for cell in competition_sheet.iter_cols(min_col=2, max_col=2, min_row=5):
        for item in cell:
            item.number_format = "#,##0"
    for column in range(3, 7):
        for cell in competition_sheet.iter_cols(min_col=column, max_col=column, min_row=5):
            for item in cell:
                item.number_format = "0.0%"
    opportunity = _table_sheet(
        workbook, name="Opportunity", title="Opportunity Research Shortlist",
        subtitle="PRIORITY_RESEARCH, WATCH, and LOW_PRIORITY are investigation labels under the governed candidate policy.",
        headers=["Priority", "Classification", "Route", "Quantitative / governed basis"],
        rows=[[item["priority"], item["classification"], item["route_label"], "; ".join(item["basis"])] for item in report.opportunity["recommendations"]],
    )
    opportunity.conditional_formatting.add(
        f"B5:B{opportunity.max_row}",
        CellIsRule(operator="equal", formula=['"PRIORITY_RESEARCH"'], fill=PatternFill("solid", fgColor=_PALE_GREEN)),
    )
    risk_sheet = _table_sheet(
        workbook, name="Risks", title="Risks and Unknowns",
        subtitle="These constraints must remain visible in operator decisions.",
        headers=["Risk", "Value", "Evidence basis"],
        rows=[[item["risk"], item["value"], item["basis"]] for item in report.risks_unknowns],
    )
    for row in (5, 6):
        risk_sheet.cell(row, 2).number_format = "0.0%"
    review = _table_sheet(
        workbook, name="Business Value Review", title="Business Value Review (Human Only)",
        subtitle="Human verdict and notes are intentionally blank. Technical acceptance must not infer business usefulness.",
        headers=["Question", "Human verdict", "Notes"],
        rows=[[item["question"], None, None] for item in report.business_value_review],
    )
    for row in range(5, review.max_row + 1):
        review.cell(row, 2).value = None
        review.cell(row, 3).value = None
        review.cell(row, 2).fill = PatternFill("solid", fgColor=_PALE_AMBER)
        review.cell(row, 3).fill = PatternFill("solid", fgColor=_PALE_AMBER)
    review.column_dimensions["A"].width = 72
    review.column_dimensions["B"].width = 20
    review.column_dimensions["C"].width = 48

    _table_sheet(
        workbook, name="Audit", title="Audit and Provenance",
        subtitle="Identifiers and fingerprints bind the report to Canonical evidence and every downstream result.",
        headers=["Field", "Value"],
        rows=[
            ["report_id", report.report_id],
            ["report_semantic_fingerprint", report.semantic_fingerprint],
            ["contract_version", report.contract_version],
            ["category", report.category],
            ["marketplace", report.marketplace],
            ["generated_at", report.generated_at],
            *[[key, value] for key, value in sorted(report.source_lineage.items())],
            ["limitations", list(report.limitations)],
        ],
    )
    expected = (
        "Executive Summary", "Market Overview", "Route Landscape", "Route Comparison",
        "Buyer Needs", "Competition", "Opportunity", "Risks",
        "Business Value Review", "Audit",
    )
    if tuple(workbook.sheetnames) != expected:
        raise RuntimeError("integrated operator workbook sheet contract drifted")
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.name}.tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(destination)


def workbook_semantic_sha256(path: Path) -> str:
    workbook = load_workbook(path, data_only=False, read_only=True)
    content = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            while values and values[-1] is None:
                values.pop()
            rows.append(values)
        while rows and not rows[-1]:
            rows.pop()
        content.append({"sheet": sheet.title, "rows": rows})
    workbook.close()
    return sha256(canonical_json(content).encode("utf-8")).hexdigest()


class IntegratedMarketReportDelivery:
    def deliver(
        self, report: IntegratedMarketReport, output_directory: Path
    ) -> IntegratedMarketReportDeliveryResult:
        if not isinstance(report, IntegratedMarketReport):
            raise TypeError("report must be an IntegratedMarketReport")
        output_directory.mkdir(parents=True, exist_ok=True)
        json_path = output_directory / MARKET_REPORT_JSON_FILENAME
        markdown_path = output_directory / OPERATOR_MARKDOWN_FILENAME
        xlsx_path = output_directory / OPERATOR_XLSX_FILENAME
        review_path = output_directory / BUSINESS_VALUE_REVIEW_FILENAME
        material = (
            (json_path, report.to_json(indent=2) + "\n"),
            (markdown_path, render_operator_markdown(report)),
            (review_path, render_business_value_review(report)),
        )
        for destination, content in material:
            temporary = destination.with_name(f".{destination.name}.tmp")
            temporary.write_text(content, encoding="utf-8", newline="\n")
            temporary.replace(destination)
        render_operator_xlsx(report, xlsx_path)
        xlsx_bytes = xlsx_path.read_bytes()
        return IntegratedMarketReportDeliveryResult(
            source_report_id=report.report_id,
            delivery_version=INTEGRATED_MARKET_REPORT_DELIVERY_VERSION,
            json_path=json_path,
            markdown_path=markdown_path,
            xlsx_path=xlsx_path,
            business_value_review_path=review_path,
            json_sha256=sha256(json_path.read_bytes()).hexdigest(),
            markdown_sha256=sha256(markdown_path.read_bytes()).hexdigest(),
            xlsx_sha256=sha256(xlsx_bytes).hexdigest(),
            xlsx_package_content_sha256=ooxml_package_content_sha256(xlsx_bytes),
            xlsx_semantic_sha256=workbook_semantic_sha256(xlsx_path),
            business_value_review_sha256=sha256(review_path.read_bytes()).hexdigest(),
        )


__all__ = (
    "BUSINESS_VALUE_REVIEW_FILENAME",
    "INTEGRATED_MARKET_REPORT_DELIVERY_VERSION",
    "MARKET_REPORT_JSON_FILENAME",
    "OPERATOR_MARKDOWN_FILENAME",
    "OPERATOR_XLSX_FILENAME",
    "IntegratedMarketReportDelivery",
    "IntegratedMarketReportDeliveryResult",
    "render_business_value_review",
    "render_operator_markdown",
    "render_operator_xlsx",
    "workbook_semantic_sha256",
)

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from amazon_product_intelligence.market_report.integrated import (
    IntegratedMarketReportDelivery,
    run_final_market_intelligence,
)
from amazon_product_intelligence.competition_intelligence import (
    CompetitionIntelligenceBuilderV0_1,
    CompetitionIntelligenceRequest,
)
from amazon_product_intelligence.contracts import (
    NormalizationStatus,
    PresenceStatus,
    SemanticStatus,
)
from amazon_product_intelligence.route_discovery_input import (
    ROUTE_INPUT_FIELD_MAPPINGS,
    RouteDiscoveryInputContext,
    build_route_discovery_input,
)
from amazon_product_intelligence.route_discovery_v2 import load_route_discovery_v2_config
from amazon_product_intelligence.semantic_engine_v2 import load_category_semantic_profile
from amazon_product_intelligence.sellersprite_import import (
    ImportValueStatus,
    build_sellersprite_canonical_evidence,
)
from amazon_product_intelligence.sellersprite_import.canonical_adapter import _BINDINGS
from tests.test_route_discovery_input_v1 import _bundles


ROOT = Path(__file__).resolve().parents[1]


def _observation(asin: str, kind: str, name: str, value, **extra):
    return {
        "provider": "provider_fixture",
        "asin": asin,
        "kind": kind,
        "name": name,
        "value": value,
        **extra,
    }


def _canonical_bundles():
    observations = []
    architectures = (
        ("Adhesive", "Wall", 120),
        ("Suction Cup", "Wall", 80),
        ("Hanging", "Hook", 55),
    )
    index = 0
    for installation, attachment, baseline_sales in architectures:
        for sequence in range(3):
            index += 1
            asin = f"B0SP43A{index:03d}"
            title = f"{installation} shower caddy organizer"
            facts = (
                ("asin", asin, "IDENTITY_RELATED"),
                ("title", title, "DESCRIPTION"),
                ("brand", f"Brand {sequence + 1}", "OTHER"),
                ("seller", f"Seller {sequence + 1}", "OTHER"),
                ("category", "shower caddies", "OTHER"),
                ("installation", installation, "ATTRIBUTE"),
                ("attachment", attachment, "ATTRIBUTE"),
            )
            observations.extend(
                _observation(
                    asin, "PRODUCT_FACT", name, value, fact_group=fact_group
                )
                for name, value, fact_group in facts
            )
            metrics = (
                ("estimated_monthly_sales", baseline_sales + sequence, "COUNT", "units"),
                ("price", 20.0 + index, "CURRENCY", "USD"),
                ("review_count", 100 * index, "COUNT", "reviews"),
                ("rating", 4.0 + sequence / 10, "RATING", "stars_5"),
                ("listing_age_days", 30 * index, "COUNT", "days"),
                ("bsr", 1000 * index, "RANK", "rank"),
                ("subcategory_bsr", 100 * index, "RANK", "rank"),
                ("sales_mom_growth", 0.02 * sequence, "RATIO", "ratio"),
                ("sales_yoy_growth", 0.03 * sequence, "RATIO", "ratio"),
            )
            if index != 1:
                metrics += (("estimated_monthly_revenue", (baseline_sales + sequence) * (20.0 + index), "CURRENCY", "USD"),)
            observations.extend(
                _observation(
                    asin,
                    "METRIC",
                    name,
                    value,
                    estimate=True,
                    unit={"dimension": dimension, "unit_code": code, "unit_system": "DOMAIN"},
                )
                for name, value, dimension, code in metrics
            )
    return _bundles({"observations": observations})


def _run():
    bundles = _canonical_bundles()
    return run_final_market_intelligence(
        canonical_bundles=bundles,
        input_context=RouteDiscoveryInputContext(
            marketplace="US",
            category="shower caddies",
            imported_at="2026-08-26T00:00:00Z",
            normalization_run_id="sp-043a-fixture",
            normalized_at="2026-08-26T00:00:00Z",
            observed_date="2026-08-26",
        ),
        semantic_profile=load_category_semantic_profile(
            ROOT / "config/category_semantic_profiles/shower_caddies.v1_1.json"
        ),
        route_config=load_route_discovery_v2_config(
            ROOT / "config/route_discovery_v2/shower_caddies.v2.json"
        ),
        generated_at="2026-08-26T00:00:00Z",
    )


def test_route_input_v11_maps_report_metrics_without_provider_selection() -> None:
    mappings = {(item.observation_kind, item.canonical_name): item for item in ROUTE_INPUT_FIELD_MAPPINGS}
    for key in (
        ("METRIC", "estimated_monthly_revenue"),
        ("METRIC", "sales_mom_growth"),
        ("METRIC", "sales_yoy_growth"),
        ("METRIC", "listing_age_days"),
        ("METRIC", "subcategory_bsr"),
        ("PRODUCT_FACT", "seller"),
    ):
        assert key in mappings
    assert all("provider" not in item.canonical_field for item in mappings.values())


def test_provider_adapter_normalizes_values_and_fail_closes_parent_self_loop() -> None:
    context = RouteDiscoveryInputContext(
        marketplace="US", category="shower caddies",
        imported_at="2026-08-26T00:00:00Z",
        normalization_run_id="sp-043a-adapter-fixture",
        normalized_at="2026-08-26T00:00:00Z",
        observed_date="2026-08-26",
    )
    dataset = build_route_discovery_input(_canonical_bundles(), context=context).dataset
    record = dataset.records[0]
    parent_header = next(item.header for item in _BINDINGS if item.name == "parent_product_relationship")
    fields = tuple(
        replace(
            field,
            value=record.asin,
            import_status=ImportValueStatus.NORMALIZED,
            presence_status=PresenceStatus.PRESENT,
            normalization_status=NormalizationStatus.NORMALIZED,
            semantic_status=SemanticStatus.CONFIRMED,
        ) if field.header == parent_header else field
        for field in record.fields
    )
    modified = replace(dataset, records=(replace(record, parent_asin=record.asin, fields=fields), *dataset.records[1:]))
    bundle = build_sellersprite_canonical_evidence(modified)
    revenue = next(
        item for item in bundle.observations
        if getattr(item, "metric", None) == "estimated_monthly_revenue"
        and item.value.presence_status is PresenceStatus.PRESENT
    )
    assert revenue.value.normalized_value == revenue.value.raw_value
    parent = next(
        item for item in bundle.observations
        if getattr(item, "dimension", None) == "parent_product_relationship"
        and item.subject.subject_id.endswith(record.asin)
    )
    assert parent.value.raw_value == record.asin
    assert parent.value.semantic_status is SemanticStatus.SEMANTICS_UNCONFIRMED
    assert "SELF_REFERENTIAL_PARENT_ASIN" in {item.issue_code for item in bundle.quality_issues}
    CompetitionIntelligenceBuilderV0_1().build(
        CompetitionIntelligenceRequest(canonical_bundles=(bundle,))
    )


def test_provider_neutral_pipeline_is_route_aware_and_preserves_missingness() -> None:
    result = _run()
    report = result.market_report
    assert result.route_input.dataset.source_kind == "CANONICAL_PROVIDER_OBSERVATIONS"
    assert report.market_overview["coverage_counts"]["assigned"] == sum(
        item["products"] for item in report.route_landscape
    )
    assert len(report.route_landscape) >= 2
    assert all(item["route_definition"] for item in report.route_landscape)
    revenue = report.market_overview["monthly_revenue"]
    assert revenue["availability"] == "PARTIAL"
    assert revenue["unknown_count"] == 1
    assert revenue["missing_values_are_zero"] is False
    assert report.buyer_needs["availability"] == "UNAVAILABLE"
    assert report.buyer_needs["themes"] == []
    assert all(
        item["human_verdict"] is None and item["notes"] is None
        for item in report.business_value_review
    )
    assert result.route_result.upstream_semantic_result_id == result.semantic_result.result_id
    assert report.source_lineage["route_result_id"] == result.route_result.result_id


def test_delivery_is_deterministic_blank_and_listing_safe(tmp_path: Path) -> None:
    first = _run()
    second = _run()
    assert first.semantic_result.to_json() == second.semantic_result.to_json()
    assert first.route_result.to_json() == second.route_result.to_json()
    assert first.market_report.to_json() == second.market_report.to_json()
    delivery = IntegratedMarketReportDelivery()
    output_a = delivery.deliver(first.market_report, tmp_path / "a")
    output_b = delivery.deliver(second.market_report, tmp_path / "b")
    assert output_a.json_sha256 == output_b.json_sha256
    assert output_a.markdown_sha256 == output_b.markdown_sha256
    assert output_a.xlsx_semantic_sha256 == output_b.xlsx_semantic_sha256
    assert output_a.xlsx_package_content_sha256 == output_b.xlsx_package_content_sha256
    report_text = output_a.markdown_path.read_text(encoding="utf-8")
    assert "B0SP43A" not in report_text
    assert "|  |  |" in report_text
    workbook = load_workbook(output_a.xlsx_path, data_only=False)
    assert tuple(workbook.sheetnames) == (
        "Executive Summary", "Market Overview", "Route Landscape", "Route Comparison",
        "Buyer Needs", "Competition", "Opportunity", "Risks",
        "Business Value Review", "Audit",
    )
    review = workbook["Business Value Review"]
    assert all(review.cell(row, 2).value is None for row in range(5, 13))
    assert all(review.cell(row, 3).value is None for row in range(5, 13))
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_report_core_contains_no_provider_or_category_specific_branch() -> None:
    core = ROOT / "src/amazon_product_intelligence/market_report/integrated"
    text = "\n".join(path.read_text(encoding="utf-8").casefold() for path in core.glob("*.py"))
    assert "sellersprite" not in text
    assert "sorftime" not in text
    assert "xiyou" not in text
    assert "shower cadd" not in text
    assert "air fryer" not in text
